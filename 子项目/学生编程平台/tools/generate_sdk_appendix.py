from __future__ import annotations

import ast
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_PATH = PROJECT_ROOT / "src" / "aibox" / "__init__.py"
OUTPUT_PATH = Path.home() / "Desktop" / "人工智能实验箱学生编程平台_aibox封装代码附录.docx"


API_SECTIONS = [
    (
        "Arm 机械臂",
        "隐藏串口、舵机协议、动作序列和安全停止逻辑。学生只使用预设动作或安全范围内的舵机微调。",
        [
            ("arm = Arm()", "创建机械臂教学对象。创建对象时不直接操作硬件。", "返回 Arm 对象", "当前模拟封装已实现"),
            ("arm.reset()", "机械臂复位，回到安全初始状态。", "返回 None", "当前模拟封装已实现"),
            ("arm.pose(\"直立\")", "执行预设姿态。支持“直立”和“放平”。", "name: str；返回 None", "只允许白名单姿态"),
            ("arm.open()", "张开夹爪。", "返回 None", "当前模拟封装已实现"),
            ("arm.half_open()", "夹爪移动到半开位置。", "返回 None", "当前模拟封装已实现"),
            ("arm.close()", "闭合夹爪。", "返回 None", "当前模拟封装已实现"),
            ("arm.action(\"抓取\")", "执行整套动作组。支持抓取、搬运、左转移、右转移。", "name: str；返回 None", "动作名必须在白名单中"),
            ("arm.command(command)", "执行语音实验识别到的固定机械臂命令。", "command: str；返回 None", "仅支持现有功能演示的10个命令"),
            ("arm.adjust(1, 30)", "对1至6号舵机进行小步微调。", "joint: 1~6；delta: -100~100", "调整后的PWM仍须处于800~2200"),
            ("arm.servo(1, 1500)", "设置单个舵机的安全PWM目标值。", "joint: 1~6；pwm: 800~2200", "越界时拒绝执行并抛出SafetyError"),
            ("arm.stop()", "立即请求停止当前机械臂动作。", "返回 None", "停止优先于普通动作"),
        ],
    ),
    (
        "Camera 摄像头",
        "学生只学习调用实时画面，不学习设备编号、分辨率、拍照文件和V4L2参数。",
        [
            ("camera = Camera()", "创建摄像头教学对象。", "返回 Camera 对象", "当前模拟封装已实现"),
            ("camera.show()", "自动申请摄像头，在右侧效果区显示实时画面。", "返回 None", "真实设备由Device Service独占管理"),
            ("camera.close()", "停止实时画面并释放摄像头。", "返回 None", "平台停止时也会自动释放"),
        ],
    ),
    (
        "Vision 视觉识别",
        "隐藏模型加载、RKNN/ONNX推理、图像预处理和结果绘制。第一版不提供通用目标检测接口。",
        [
            ("Vision().color()", "识别画面中的颜色，并在右侧显示结果。", "返回颜色中文名称 str；无结果返回空字符串", "当前模拟值为“红色”"),
            ("Vision().fruit()", "识别画面中的水果。", "返回水果中文名称 str；无结果返回空字符串", "当前模拟值为“苹果”"),
            ("Vision().classify()", "对画面中的教学物体进行分类。", "返回类别中文名称 str；无结果返回空字符串", "当前模拟值为“积木”"),
            ("Vision().shape()", "识别轮廓形状。", "返回形状中文名称 str；无结果返回空字符串", "当前模拟值为“圆形”"),
        ],
    ),
    (
        "Face 人脸检测",
        "只检测人脸数量，不进行身份识别，也不保存人脸信息。",
        [
            ("count = Face().detect()", "检测当前画面中的人脸数量。", "返回 int；无人脸返回0", "当前模拟值为1"),
        ],
    ),
    (
        "Hand 手掌与手势",
        "隐藏手部模型、关键点处理和云台跟踪控制。",
        [
            ("Hand().detect()", "判断是否检测到手掌。", "返回 bool", "当前模拟值为True"),
            ("Hand().gesture()", "识别石头、剪刀、布等教学手势。", "返回手势中文名称 str", "当前模拟值为“剪刀”"),
            ("Hand().track()", "启动手掌跟踪，右侧持续显示效果并控制云台。", "返回 None", "属于观察型功能，点击停止后结束"),
            ("Hand().stop()", "停止手掌跟踪和云台动作。", "返回 None", "可重复调用"),
        ],
    ),
    (
        "Voice 离线语音",
        "隐藏麦克风、ALSA、ASR、TTS模型和声卡选择，所有教学功能离线运行。",
        [
            ("text = Voice().listen()", "通过麦克风听取中文语音并转换成文字。", "timeout默认8秒；返回 str", "真机阶段复用功能演示ASR"),
            ("Voice().speak(\"你好\")", "把代码中的文字转换成语音并通过实验箱声卡播放。", "text为1~200个字符；返回 None", "空文字或超长文字会被拒绝"),
            ("Voice().stop()", "停止当前录音、识别或播放任务。", "返回 None", "可重复调用"),
        ],
    ),
    (
        "AI 本地对话",
        "隐藏本地大模型加载、提示词和生成参数，断网时仍可运行。",
        [
            ("answer = AI().chat(\"什么是人工智能？\")", "向本地离线AI提问并获得中文回答。", "text: 1~500字；返回 str", "不访问互联网"),
            ("AI().stop()", "中止当前本地AI回答。", "返回 None", "可重复调用"),
        ],
    ),
    (
        "Lidar 雷达",
        "隐藏思岚A1M8串口、波特率、协议解析、角度补偿和扫描线程。教材名称固定使用Lidar。",
        [
            ("lidar = Lidar()", "创建雷达教学对象。", "返回 Lidar 对象", "当前模拟封装已实现"),
            ("lidar.start()", "显式启动A1M8扫描。通常由show()自动完成。", "返回 None", "目标接口；真机Device Service阶段实现"),
            ("lidar.show()", "在右侧效果区显示雷达扫描图。", "返回 None", "属于观察型功能"),
            ("distance = lidar.get_distance()", "读取课程固定区域内的代表距离。", "返回毫米 float；无目标返回-1.0", "当前模拟值为860.0"),
            ("changed = lidar.detect()", "判断固定检测区域是否发生物体变化。", "返回 bool", "不向学生暴露原始扫描点"),
            ("lidar.stop()", "停止扫描并释放雷达串口。", "返回 None", "可重复调用"),
        ],
    ),
]


COMBINED_EXAMPLES = [
    (
        "语音控制机械臂",
        """from aibox import Arm, Voice

command = Voice().listen()
Arm().command(command)
print("执行语音指令：", command)""",
        "把离线语音识别结果交给机械臂固定命令表。支持直立、放平、抓取、搬运、复位、停止、张开、闭合、右转移和左转移。",
    ),
    (
        "机械臂打开后说话",
        """from aibox import Arm, Voice

Arm().open()
Voice().speak("机械臂已经张开")""",
        "先张开机械臂夹爪，再通过离线TTS和声卡播报状态。",
    ),
    (
        "显示摄像头画面",
        """from aibox import Camera

camera = Camera()
camera.show()""",
        "只调用实时画面。学生不需要设置摄像头编号、分辨率或底层驱动。",
    ),
    (
        "雷达扫描与测距",
        """from aibox import Lidar

lidar = Lidar()
lidar.show()
distance = lidar.get_distance()
print("检测距离：", distance, "毫米")""",
        "显示A1M8扫描图，并读取固定教学区域的代表距离。",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_run_font(run, name: str, size: float, *, bold: bool = False, color: str | None = None) -> None:
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    for style_name, size, color in (
        ("Title", 22, "173A2B"),
        ("Heading 1", 16, "173A2B"),
        ("Heading 2", 13, "256047"),
        ("Heading 3", 11, "2E493C"),
    ):
        style = document.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)


def add_code_paragraph(container, code: str, *, compact: bool = False) -> None:
    paragraph = container.add_paragraph() if hasattr(container, "add_paragraph") else container
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.left_indent = Cm(0.2)
    run = paragraph.add_run(code)
    set_run_font(run, "Consolas", 8.5 if compact else 9, color="173A2B")


def add_api_table(document: Document, rows: list[tuple[str, str, str, str]]) -> None:
    table = document.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headings = ("学生调用代码", "代码作用", "参数与返回值", "状态/安全说明")
    for index, heading in enumerate(headings):
        cell = table.rows[0].cells[index]
        cell.text = heading
        set_cell_shading(cell, "DCECE5")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, "Microsoft YaHei", 9, bold=True, color="173A2B")
    set_repeat_table_header(table.rows[0])

    for code, purpose, result, safety in rows:
        cells = table.add_row().cells
        add_code_paragraph(cells[0], code, compact=True)
        cells[1].text = purpose
        cells[2].text = result
        cells[3].text = safety
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(1)
                for run in paragraph.runs:
                    if cell is not cells[0]:
                        set_run_font(run, "Microsoft YaHei", 8.5)

    widths = (Cm(5.0), Cm(7.0), Cm(4.4), Cm(5.3))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width


def add_source_listing(document: Document, source: str) -> None:
    for index, line in enumerate(source.splitlines(), start=1):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1
        number = paragraph.add_run(f"{index:>3}  ")
        set_run_font(number, "Consolas", 7.5, color="789085")
        code = paragraph.add_run(line or " ")
        set_run_font(code, "Consolas", 7.5, color="17221D")
        if index % 2 == 0:
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F4F6F4")
            paragraph._p.get_or_add_pPr().append(shading)


def validate_document_scope(source: str) -> None:
    tree = ast.parse(source)
    public = {
        node.name: {
            child.name
            for child in node.body
            if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef)) and not child.name.startswith("_")
        }
        for node in tree.body
        if isinstance(node, ast.ClassDef) and node.name in {"Arm", "Camera", "Vision", "Face", "Hand", "Voice", "AI", "Lidar"}
    }
    documented = {
        "Arm": {"reset", "pose", "open", "half_open", "close", "action", "command", "adjust", "servo", "stop"},
        "Camera": {"show", "close"},
        "Vision": {"color", "fruit", "classify", "shape"},
        "Face": {"detect"},
        "Hand": {"detect", "gesture", "track", "stop"},
        "Voice": {"listen", "speak", "stop"},
        "AI": {"chat", "stop"},
        "Lidar": {"show", "get_distance", "detect", "stop"},
    }
    if public != documented:
        raise RuntimeError(f"附录接口清单与当前封装源码不一致：{public!r}")


def build_document() -> Path:
    source = SOURCE_PATH.read_text(encoding="utf-8")
    validate_document_scope(source)

    document = Document()
    configure_document(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("附录：aibox SDK 全部封装代码与作用说明")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"人工智能实验箱学生编程平台 · 版本 1.0-draft · {date.today().isoformat()}")
    set_run_font(run, "Microsoft YaHei", 10, color="607269")

    notice = document.add_table(rows=1, cols=1)
    notice.style = "Table Grid"
    notice.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = notice.cell(0, 0)
    set_cell_shading(cell, "FFF4D6")
    cell.text = (
        "重要说明：本附录中的学生API是教材与软件的目标合同。当前原型使用模拟封装，"
        "不会连接真实硬件；RK3588正式版将保持相同学生调用方式，在Device Service中替换为"
        "真实摄像头、机械臂、雷达、语音和本地AI实现。学生不得直接调用串口、USB、GPIO、"
        "ALSA、V4L2、模型路径或底层驱动。"
    )
    for run in cell.paragraphs[0].runs:
        set_run_font(run, "Microsoft YaHei", 9.5, bold=True, color="684E12")

    document.add_heading("一、统一导入方式", level=1)
    add_code_paragraph(document, "from aibox import Arm, Camera, Vision, Face, Hand, Voice, AI, Lidar")
    paragraph = document.add_paragraph(
        "学生只需要导入 aibox 中的教学对象。所有设备通信、资源互斥、线程、模型和错误翻译由平台完成。"
    )
    paragraph.paragraph_format.space_after = Pt(8)

    document.add_heading("二、全部学生封装API", level=1)
    for title_text, description, rows in API_SECTIONS:
        document.add_heading(title_text, level=2)
        document.add_paragraph(description)
        add_api_table(document, rows)
        document.add_paragraph()

    document.add_heading("三、组合实验代码", level=1)
    for title_text, code, purpose in COMBINED_EXAMPLES:
        document.add_heading(title_text, level=2)
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        add_code_paragraph(table.cell(0, 0), code)
        table.cell(0, 1).text = purpose
        set_cell_shading(table.cell(0, 0), "F1F5F2")
        for run in table.cell(0, 1).paragraphs[0].runs:
            set_run_font(run, "Microsoft YaHei", 9)

    document.add_heading("四、异常与安全限制", level=1)
    safety_rows = [
        ("SafetyError", "舵机编号、PWM、微调步长、朗读文字长度或动作名称超出安全范围。"),
        ("DeviceNotFoundError", "正式版未找到摄像头、机械臂、雷达、麦克风或其他设备。"),
        ("DeviceBusyError", "设备正在被功能演示或其他实验占用。"),
        ("OperationTimeoutError", "设备操作或模型运行超过限定时间。"),
        ("ModelError", "离线视觉、语音或AI模型无法加载。"),
    ]
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "异常名称"
    table.rows[0].cells[1].text = "作用"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "DCECE5")
        for run in cell.paragraphs[0].runs:
            set_run_font(run, "Microsoft YaHei", 9, bold=True)
    for name, purpose in safety_rows:
        cells = table.add_row().cells
        add_code_paragraph(cells[0], name, compact=True)
        cells[1].text = purpose

    document.add_heading("五、封装层与真实硬件的关系", level=1)
    architecture = (
        "学生代码 → aibox SDK → 本地受控RPC → Device Service → 硬件适配器 → 真实设备"
    )
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(architecture)
    set_run_font(run, "Microsoft YaHei", 11, bold=True, color="176B4A")
    document.add_paragraph(
        "正式版中，aibox SDK不得直接打开 /dev 设备节点。机械臂、摄像头、雷达、麦克风、"
        "声卡和NPU都由Device Service统一管理。现有“功能演示”和学生平台必须共享资源锁，"
        "防止同一设备被两个软件同时占用。"
    )

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("六、当前模拟封装源码全文", level=1)
    document.add_paragraph(
        "以下代码来自当前开发原型的 src/aibox/__init__.py。它用于验证教材API、课程模板和界面流程，"
        "通过事件模拟设备效果；它不是最终的RK3588真实硬件驱动。"
    )
    add_source_listing(document, source)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document())
